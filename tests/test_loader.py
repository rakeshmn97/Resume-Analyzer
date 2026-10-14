from pathlib import Path

import pytest

from core.loader import guess_candidate_name, load_resume

ROOT = Path(__file__).resolve().parents[1]
SAMPLES = ROOT / "data" / "samples"


def test_load_sample_txt_and_guess_name():
    path = SAMPLES / "priya_sharma.txt"
    text = load_resume(str(path))
    assert "PyTorch" in text
    assert guess_candidate_name(text, path.name) == "Priya Sharma"


def test_load_docx(tmp_path: Path):
    from docx import Document

    document = Document()
    document.add_paragraph("Rahul Verma")
    document.add_paragraph("Skills:")
    document.add_paragraph("Python and SQL")
    path = tmp_path / "rahul.docx"
    document.save(path)
    text = load_resume(str(path))
    assert "Rahul Verma" in text
    assert "Python and SQL" in text


def test_reject_unknown_type(tmp_path: Path):
    path = tmp_path / "notes.md"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(ValueError):
        load_resume(str(path))
